from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from accounts.models import AdminProfile, CustomUser, LecturerProfile, StudentProfile
from courses.models import AcademicSession, Course, CourseAllocation, Enrollment, Programme
from .grading import GradingEngine
from .models import Result, SemesterGPA, SemesterResultBatch, CumulativeGPA
from .services import get_batch_progress
from .signals import recompute_gpa
from .views import (
	ApproveBatchView,
	BatchListView,
	LecturerResultExcelUploadView,
	LecturerResultUploadView,
	ScoreSheetTemplateView,
	StudentGPAView,
	StudentResultsView,
)


class ResultWorkflowTests(TestCase):
	def setUp(self):
		self.factory = APIRequestFactory()
		self.session = AcademicSession.objects.create(
			name="2099/2100", start_date="2099-09-01", end_date="2100-08-31", is_current=True
		)
		self.programme = Programme.objects.create(name="Computer Science", programme_type="MSC")
		self.admin = CustomUser.objects.create_user(
			email="admin@test.com", password="password123", role="ADMIN", first_name="PG", last_name="Coordinator"
		)
		AdminProfile.objects.create(user=self.admin, admin_role="PG_COORDINATOR")
		self.lecturers = []
		for index in range(3):
			user = CustomUser.objects.create_user(
				email=f"lecturer{index}@test.com", password="password123", role="LECTURER",
				first_name=f"Lecturer{index}", last_name="Test",
			)
			self.lecturers.append(LecturerProfile.objects.create(
				user=user, staff_id=f"STAFF-{index}", status="ACTIVE"
			))
		student_user = CustomUser.objects.create_user(
			email="student@test.com", password="password123", role="STUDENT",
			first_name="Student", last_name="Test",
		)
		self.student = StudentProfile.objects.create(
			user=student_user, programme="MSC", matric_number="MSC-001", admission_year=2099, status="ACTIVE"
		)
		self.allocations = []
		for index, lecturer in enumerate(self.lecturers, start=1):
			course = Course.objects.create(
				code=f"CSC9{index:02d}", title=f"Test Course {index}", credit_units=3,
				semester="FIRST", programme=self.programme,
			)
			allocation = CourseAllocation.objects.create(
				lecturer=lecturer, course=course, session=self.session, semester="FIRST"
			)
			self.allocations.append(allocation)
			Enrollment.objects.create(student=self.student, allocation=allocation, status="APPROVED")
		self.batch = SemesterResultBatch.objects.create(
			session=self.session, semester="FIRST", programme_type="MSC"
		)

	def upload_json(self, lecturer, enrollment, score="70"):
		request = self.factory.post(
			"/api/results/upload/",
			{"batch": self.batch.id, "results": [{"enrollment": enrollment.id, "score": score}]},
			format="json",
		)
		force_authenticate(request, user=lecturer.user)
		return LecturerResultUploadView.as_view()(request)

	def downloaded_template(self, allocation):
		request = self.factory.get(f"/api/results/template/{allocation.id}/")
		force_authenticate(request, user=allocation.lecturer.user)
		response = ScoreSheetTemplateView.as_view()(request, allocation_pk=allocation.id)
		self.assertEqual(response.status_code, 200)
		return load_workbook(BytesIO(response.content))

	def upload_workbook(self, lecturer, allocation, workbook):
		file = BytesIO()
		workbook.save(file)
		file.seek(0)
		file.name = "CSC_TEST_SCORE_SHEET.xlsx"
		request = self.factory.post(
			f"/api/results/upload-excel/{allocation.id}/",
			{"file": file},
			format="multipart",
		)
		force_authenticate(request, user=lecturer.user)
		return LecturerResultExcelUploadView.as_view()(request, allocation_pk=allocation.id)

	def approve(self):
		request = self.factory.post(
			f"/api/results/batches/{self.batch.id}/approve/", {"action": "approve"}, format="json"
		)
		force_authenticate(request, user=self.admin)
		return ApproveBatchView.as_view()(request, pk=self.batch.id)

	def test_multiple_lecturers_share_one_pending_batch(self):
		for lecturer, allocation in zip(self.lecturers, self.allocations):
			response = self.upload_json(lecturer, allocation.enrollments.get())
			self.assertEqual(response.status_code, 200)
		self.assertEqual(Result.objects.filter(batch=self.batch).count(), 3)
		self.assertEqual(SemesterResultBatch.objects.count(), 1)

	def test_coordinator_batch_list_exposes_all_course_progress(self):
		self.upload_json(self.lecturers[0], self.allocations[0].enrollments.get())
		request = self.factory.get("/api/results/batches/")
		force_authenticate(request, user=self.admin)
		response = BatchListView.as_view()(request)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(len(response.data), 1)
		progress = response.data[0]["progress"]
		self.assertEqual(progress["expected_course_count"], 3)
		self.assertEqual(progress["submitted_course_count"], 1)
		self.assertEqual(progress["missing_course_count"], 2)
		first_course = next(
			course for course in progress["courses"]
			if course["allocation"] == self.allocations[0].id
		)
		self.assertEqual(first_course["status"], "COMPLETE")
		self.assertEqual(first_course["lecturer"], "Lecturer0 Test")

	def test_excel_uploads_from_multiple_lecturers_create_one_shared_batch(self):
		self.batch.delete()
		batch_ids = set()
		for lecturer, allocation in zip(self.lecturers, self.allocations):
			workbook = self.downloaded_template(allocation)
			workbook.active.cell(row=10, column=4, value=25)
			workbook.active.cell(row=10, column=5, value=50)
			response = self.upload_workbook(lecturer, allocation, workbook)
			self.assertEqual(response.status_code, 200)
			batch_ids.add(response.data["batch"])
		self.assertEqual(len(batch_ids), 1)
		self.assertEqual(SemesterResultBatch.objects.count(), 1)
		self.assertEqual(Result.objects.count(), 3)

	def test_incomplete_batch_cannot_be_approved(self):
		self.upload_json(self.lecturers[0], self.allocations[0].enrollments.get())
		response = self.approve()
		self.assertEqual(response.status_code, 400)
		self.batch.refresh_from_db()
		self.assertEqual(self.batch.status, SemesterResultBatch.Status.PENDING)
		self.assertEqual(response.data["progress"]["missing_course_count"], 2)
		self.assertEqual(
			{item["status"] for item in response.data["progress"]["courses"]},
			{"COMPLETE", "NOT_UPLOADED"},
		)
		self.assertEqual(
			{item["course_code"] for item in response.data["outstanding_courses"]},
			{self.allocations[1].course.code, self.allocations[2].course.code},
		)

	def test_course_progress_counts_all_approved_enrollments_and_reports_missing_student(self):
		second_user = CustomUser.objects.create_user(
			email="student2@test.com", password="test-password-123", role="STUDENT",
			first_name="Second", last_name="Student",
		)
		second_student = StudentProfile.objects.create(
			user=second_user, programme="MSC", matric_number="MSC-002",
			admission_year=2099, status="ACTIVE",
		)
		Enrollment.objects.create(
			student=second_student, allocation=self.allocations[0], status="APPROVED"
		)
		self.upload_json(self.lecturers[0], self.allocations[0].enrollments.get(student=self.student))

		progress = get_batch_progress(self.batch)
		first_course = next(
			course for course in progress["courses"]
			if course["allocation"] == self.allocations[0].id
		)
		self.assertEqual(first_course["status"], "INCOMPLETE")
		self.assertEqual(first_course["expected_result_count"], 2)
		self.assertEqual(first_course["result_count"], 1)
		self.assertEqual(first_course["missing_students"], [{
			"matric_number": "MSC-002",
			"student_name": "Second Student",
		}])
		self.assertEqual(progress["missing_course_count"], 3)

	def test_batch_progress_is_scoped_to_programme_session_and_semester(self):
		other_session = AcademicSession.objects.create(
			name="2100/2101", start_date="2100-09-01", end_date="2101-08-31"
		)
		other_programme = Programme.objects.create(name="Data Science", programme_type="PHD")
		cases = [
			("CSCSECOND", self.programme, self.session, "SECOND"),
			("PHDCOURSE", other_programme, self.session, "FIRST"),
			("NEXTSESSION", self.programme, other_session, "FIRST"),
		]
		for code, programme, session, semester in cases:
			course = Course.objects.create(
				code=code, title=code, credit_units=3, semester=semester,
				programme=programme,
			)
			CourseAllocation.objects.create(
				lecturer=self.lecturers[0], course=course, session=session, semester=semester
			)

		progress = get_batch_progress(self.batch)
		self.assertEqual(progress["expected_course_count"], 3)
		self.assertEqual(
			{course["course_code"] for course in progress["courses"]},
			{allocation.course.code for allocation in self.allocations},
		)

	def test_zero_enrollment_course_does_not_block_complete_batch_approval(self):
		course = Course.objects.create(
			code="CSC900", title="No Enrollments", credit_units=3,
			semester="FIRST", programme=self.programme,
		)
		CourseAllocation.objects.create(
			lecturer=self.lecturers[0], course=course, session=self.session, semester="FIRST"
		)
		for lecturer, allocation in zip(self.lecturers, self.allocations):
			self.upload_json(lecturer, allocation.enrollments.get())

		progress = get_batch_progress(self.batch)
		no_enrollment_course = next(
			item for item in progress["courses"] if item["course_code"] == "CSC900"
		)
		self.assertEqual(no_enrollment_course["status"], "NO_ENROLLMENTS")
		self.assertEqual(no_enrollment_course["expected_result_count"], 0)
		self.assertEqual(no_enrollment_course["result_count"], 0)
		self.assertEqual(progress["no_enrollment_course_count"], 1)
		self.assertEqual(progress["missing_course_count"], 0)
		self.assertEqual(self.approve().status_code, 200)

	def test_batch_with_missing_student_result_cannot_be_approved(self):
		other_user = CustomUser.objects.create_user(
			email="student2@test.com", password="test-password", role="STUDENT",
			first_name="Second", last_name="Student",
		)
		other_student = StudentProfile.objects.create(
			user=other_user, programme="MSC", matric_number="MSC-002",
			admission_year=2099, status="ACTIVE",
		)
		Enrollment.objects.create(
			student=other_student, allocation=self.allocations[0], status="APPROVED",
		)
		for lecturer, allocation in zip(self.lecturers, self.allocations):
			self.upload_json(lecturer, allocation.enrollments.get(student=self.student))

		response = self.approve()
		self.assertEqual(response.status_code, 400)
		progress = next(
			course for course in response.data["progress"]["courses"]
			if course["allocation"] == self.allocations[0].id
		)
		self.assertEqual(progress["result_count"], 1)
		self.assertEqual(progress["expected_result_count"], 2)
		self.assertFalse(progress["submitted"])
		self.batch.refresh_from_db()
		self.assertEqual(self.batch.status, SemesterResultBatch.Status.PENDING)

	def test_all_required_courses_submitted_allows_approval(self):
		for lecturer, allocation in zip(self.lecturers, self.allocations):
			self.upload_json(lecturer, allocation.enrollments.get())
		response = self.approve()
		self.assertEqual(response.status_code, 200)
		self.batch.refresh_from_db()
		self.assertEqual(self.batch.status, SemesterResultBatch.Status.APPROVED)

	def test_approved_batch_rejects_further_uploads(self):
		for lecturer, allocation in zip(self.lecturers, self.allocations):
			self.upload_json(lecturer, allocation.enrollments.get())
		self.assertEqual(self.approve().status_code, 200)
		response = self.upload_json(self.lecturers[0], self.allocations[0].enrollments.get(), score="80")
		self.assertEqual(response.status_code, 400)
		self.assertIn("approved", response.data["detail"].lower())

	def test_approved_batch_cannot_be_rejected_or_unlocked(self):
		for lecturer, allocation in zip(self.lecturers, self.allocations):
			self.upload_json(lecturer, allocation.enrollments.get())
		self.assertEqual(self.approve().status_code, 200)

		request = self.factory.post(
			f"/api/results/batches/{self.batch.id}/approve/",
			{"action": "reject", "comment": "attempted unlock"},
			format="json",
		)
		force_authenticate(request, user=self.admin)
		response = ApproveBatchView.as_view()(request, pk=self.batch.id)
		self.assertEqual(response.status_code, 400)
		self.batch.refresh_from_db()
		self.assertEqual(self.batch.status, SemesterResultBatch.Status.APPROVED)
		self.assertEqual(Result.objects.filter(batch=self.batch).count(), 3)

	def test_duplicate_upload_updates_existing_result(self):
		enrollment = self.allocations[0].enrollments.get()
		self.assertEqual(self.upload_json(self.lecturers[0], enrollment, "70").status_code, 200)
		result_id = Result.objects.get(enrollment=enrollment).id
		self.assertEqual(self.upload_json(self.lecturers[0], enrollment, "80").status_code, 200)
		self.assertEqual(Result.objects.filter(enrollment=enrollment).count(), 1)
		result = Result.objects.get(enrollment=enrollment)
		self.assertEqual(result.id, result_id)
		self.assertEqual(result.score, Decimal("80.00"))

	def test_only_approved_results_contribute_to_gpa(self):
		enrollment = self.allocations[0].enrollments.get()
		self.upload_json(self.lecturers[0], enrollment)
		self.assertFalse(SemesterGPA.objects.filter(student=self.student).exists())
		self.batch.status = SemesterResultBatch.Status.APPROVED
		self.batch.save()
		recompute_gpa(self.student, self.session, "FIRST")
		gpa = SemesterGPA.objects.get(student=self.student, session=self.session, semester="FIRST")
		self.assertEqual(gpa.gpa, Decimal("5.00"))

	def test_student_can_retrieve_approved_results(self):
		enrollment = self.allocations[0].enrollments.get()
		self.upload_json(self.lecturers[0], enrollment)
		self.batch.status = SemesterResultBatch.Status.APPROVED
		self.batch.save()
		request = self.factory.get("/api/results/mine/")
		force_authenticate(request, user=self.student.user)
		response = StudentResultsView.as_view()(request)
		self.assertEqual(response.status_code, 200)
		self.assertEqual(len(response.data), 1)

	def test_lecturer_cannot_upload_another_lecturers_allocation(self):
		enrollment = self.allocations[1].enrollments.get()
		response = self.upload_json(self.lecturers[0], enrollment)
		self.assertEqual(response.status_code, 403)
		self.assertFalse(Result.objects.filter(enrollment=enrollment).exists())

	def test_lecturer_cannot_upload_results_for_unapproved_enrollment(self):
		other_user = CustomUser.objects.create_user(
			email="pending-student@test.com", password="test-password", role="STUDENT",
			first_name="Pending", last_name="Student",
		)
		other_student = StudentProfile.objects.create(
			user=other_user, programme="MSC", matric_number="MSC-003",
			admission_year=2099, status="ACTIVE",
		)
		enrollment = Enrollment.objects.create(
			student=other_student,
			allocation=self.allocations[0],
			status="PENDING",
		)

		response = self.upload_json(self.lecturers[0], enrollment)
		self.assertEqual(response.status_code, 400)
		self.assertIn("approved enrollments", response.data["detail"])
		self.assertFalse(Result.objects.filter(enrollment=enrollment).exists())

	def test_student_gpa_endpoint_returns_gpa_and_cgpa_after_approval(self):
		enrollment = self.allocations[0].enrollments.get()
		self.upload_json(self.lecturers[0], enrollment)
		result = Result.objects.get(enrollment=enrollment)
		self.batch.status = SemesterResultBatch.Status.APPROVED
		self.batch.save()
		result.save()
		request = self.factory.get("/api/results/gpa/")
		force_authenticate(request, user=self.student.user)
		response = StudentGPAView.as_view()(request)
		self.assertEqual(response.status_code, 200)
		self.assertEqual(len(response.data["semester_gpas"]), 1)
		self.assertIsNotNone(response.data["cumulative"])

	def test_batch_approval_publishes_results_and_recomputes_gpa(self):
		for lecturer, allocation in zip(self.lecturers, self.allocations):
			self.upload_json(lecturer, allocation.enrollments.get())

		response = self.approve()
		self.assertEqual(response.status_code, 200)

		gpa = SemesterGPA.objects.get(
			student=self.student, session=self.session, semester="FIRST"
		)
		cgpa = CumulativeGPA.objects.get(student=self.student)
		self.assertEqual(gpa.gpa, Decimal("5.00"))
		self.assertEqual(cgpa.cgpa, Decimal("5.00"))

		request = self.factory.get("/api/results/mine/")
		force_authenticate(request, user=self.student.user)
		student_response = StudentResultsView.as_view()(request)
		self.assertEqual(student_response.status_code, 200)
		self.assertEqual(len(student_response.data), 3)

	def test_downloaded_template_with_scores_uploads_successfully(self):
		allocation = self.allocations[0]
		workbook = self.downloaded_template(allocation)
		worksheet = workbook.active
		worksheet.cell(row=10, column=4, value=25)
		worksheet.cell(row=10, column=5, value=50)
		response = self.upload_workbook(allocation.lecturer, allocation, workbook)
		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.data["uploaded"], 1)
		result = Result.objects.get(enrollment=allocation.enrollments.get())
		self.assertEqual(result.score, Decimal("75.00"))

	def test_invalid_score_is_reported(self):
		allocation = self.allocations[0]
		workbook = self.downloaded_template(allocation)
		worksheet = workbook.active
		worksheet.cell(row=10, column=4, value=80)
		worksheet.cell(row=10, column=5, value=30)
		response = self.upload_workbook(allocation.lecturer, allocation, workbook)
		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.data["uploaded"], 0)
		self.assertIn("exceeds 100", response.data["errors"][0]["detail"])

	def test_non_finite_excel_score_is_reported_without_server_error(self):
		allocation = self.allocations[0]
		workbook = self.downloaded_template(allocation)
		workbook.active.cell(row=10, column=4, value="NaN")
		workbook.active.cell(row=10, column=5, value=50)
		response = self.upload_workbook(allocation.lecturer, allocation, workbook)
		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.data["uploaded"], 0)
		self.assertIn("finite numbers", response.data["errors"][0]["detail"])

	def test_wrong_course_metadata_is_rejected(self):
		allocation = self.allocations[0]
		workbook = self.downloaded_template(allocation)
		workbook.active.cell(row=5, column=1, value="CSC999 : Wrong Course")
		workbook.active.cell(row=8, column=1, value="CSC999")
		response = self.upload_workbook(allocation.lecturer, allocation, workbook)
		self.assertEqual(response.status_code, 400)
		self.assertIn("Wrong course", response.data["detail"])

	def test_approved_batch_rejects_excel_upload(self):
		allocation = self.allocations[0]
		workbook = self.downloaded_template(allocation)
		workbook.active.cell(row=10, column=4, value=25)
		workbook.active.cell(row=10, column=5, value=50)
		self.batch.status = SemesterResultBatch.Status.APPROVED
		self.batch.save()
		response = self.upload_workbook(allocation.lecturer, allocation, workbook)
		self.assertEqual(response.status_code, 400)
		self.assertIn("approved", response.data["detail"].lower())


class GradingEngineBoundaryTests(TestCase):
	def test_course_score_boundaries(self):
		cases = [
			(39, "F", 0),
			(40, "E", 1),
			(44, "E", 1),
			(45, "D", 2),
			(49, "D", 2),
			(50, "C", 3),
			(59, "C", 3),
			(60, "B", 4),
			(69, "B", 4),
			(70, "A", 5),
			(100, "A", 5),
		]
		for score, expected_grade, expected_gp in cases:
			with self.subTest(score=score):
				self.assertEqual(GradingEngine.get_grade(score), expected_grade)
				self.assertEqual(GradingEngine.get_grade_point(score), expected_gp)

	def test_cgpa_classification_boundaries(self):
		cases = [
			(Decimal("0.99"), "Fail"),
			(Decimal("1.00"), "Pass"),
			(Decimal("1.49"), "Pass"),
			(Decimal("1.50"), "Third Class"),
			(Decimal("2.39"), "Third Class"),
			(Decimal("2.40"), "Second Class Lower"),
			(Decimal("3.49"), "Second Class Lower"),
			(Decimal("3.50"), "Second Class Upper"),
			(Decimal("4.49"), "Second Class Upper"),
			(Decimal("4.50"), "First Class"),
			(Decimal("5.00"), "First Class"),
		]
		for cgpa, expected_class in cases:
			with self.subTest(cgpa=cgpa):
				self.assertEqual(GradingEngine.classify(cgpa), expected_class)

	def test_gpa_and_cgpa_credit_weighting(self):
		# Course 1: 3 units, Score 42 -> E -> 1 gp (QP = 3)
		# Course 2: 2 units, Score 75 -> A -> 5 gp (QP = 10)
		# Total QP = 13, Total CU = 5 -> GPA = 13 / 5 = 2.60 -> Second Class Lower
		courses = [
			{"credit_units": 3, "grade_point": 1},
			{"credit_units": 2, "grade_point": 5},
		]
		gpa = GradingEngine.compute_gpa(courses)
		self.assertEqual(gpa, Decimal("2.60"))
		self.assertEqual(GradingEngine.classify(gpa), "Second Class Lower")
